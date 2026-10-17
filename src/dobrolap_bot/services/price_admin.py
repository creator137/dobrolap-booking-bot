"""Admin tool: export/import the price matrix (daily_rates) as an .xlsx file.

Lets the owner edit prices — and add brand new tariff/pet/size rows — in
Excel instead of hand-editing YAML. Export via /prices, edit the sheet
(add/change/delete rows), send the file back; the bot validates and applies
it immediately (in-memory + persisted to config/daily_rates.yaml).
"""

from __future__ import annotations

from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from dobrolap_bot.config.loader import Catalog
from dobrolap_bot.domain.enums import DogSizeClass, PetKind
from dobrolap_bot.domain.models import DailyRate

SHEET_PRICES = "Цены"
SHEET_LEGEND = "Легенда"

HEADERS = [
    "Тариф (код)",
    "Вид питомца",
    "Размер (только собаки)",
    "Молодой (щенок/котёнок)",
    "До 1 месяца",
    "Цена, ₽/сутки",
    "Цена от, ₽ (диапазон)",
    "Цена до, ₽ (диапазон)",
    "Сезон: цена, ₽/сутки",
    "Сезон: цена от, ₽",
    "Сезон: цена до, ₽",
    "Недоступно",
]

PET_KIND_LABELS: dict[PetKind, str] = {
    PetKind.DOG: "dog (собака)",
    PetKind.CAT: "cat (кошка)",
    PetKind.RABBIT: "rabbit (кролик)",
    PetKind.RAT: "rat (крыса)",
    PetKind.HAMSTER: "hamster (хомяк)",
    PetKind.BIRD: "bird (птица)",
    PetKind.GUINEA_PIG: "guinea_pig (морская свинка)",
    PetKind.OTHER: "other (другое)",
}

SIZE_LABELS: dict[DogSizeClass, str] = {
    DogSizeClass.MINIATURE: "miniature (до 10 кг)",
    DogSizeClass.MEDIUM: "medium (10–20 кг)",
    DogSizeClass.LARGE: "large (от 20 кг)",
}

_PET_KIND_ALIASES: dict[str, PetKind] = {
    "dog": PetKind.DOG,
    "собака": PetKind.DOG,
    "пёс": PetKind.DOG,
    "пес": PetKind.DOG,
    "cat": PetKind.CAT,
    "кошка": PetKind.CAT,
    "кот": PetKind.CAT,
    "rabbit": PetKind.RABBIT,
    "кролик": PetKind.RABBIT,
    "rat": PetKind.RAT,
    "крыса": PetKind.RAT,
    "hamster": PetKind.HAMSTER,
    "хомяк": PetKind.HAMSTER,
    "bird": PetKind.BIRD,
    "птица": PetKind.BIRD,
    "guinea_pig": PetKind.GUINEA_PIG,
    "морская свинка": PetKind.GUINEA_PIG,
    "свинка": PetKind.GUINEA_PIG,
    "other": PetKind.OTHER,
    "другое": PetKind.OTHER,
    "прочие животные": PetKind.OTHER,
}

_SIZE_ALIASES: dict[str, DogSizeClass] = {
    "miniature": DogSizeClass.MINIATURE,
    "миниатюрный": DogSizeClass.MINIATURE,
    "мини": DogSizeClass.MINIATURE,
    "medium": DogSizeClass.MEDIUM,
    "средний": DogSizeClass.MEDIUM,
    "large": DogSizeClass.LARGE,
    "крупный": DogSizeClass.LARGE,
}

_TRUE_WORDS = {"да", "true", "1", "yes", "истина"}
_FALSE_WORDS = {"нет", "false", "0", "no", "ложь", ""}


def _cell_str(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_bool(value: object) -> bool:
    text = _cell_str(value).lower()
    if text in _TRUE_WORDS:
        return True
    if text in _FALSE_WORDS:
        return False
    raise ValueError(f"не распознано да/нет: {value!r}")


def _parse_pet_kind(value: object) -> PetKind:
    text = _cell_str(value).lower()
    # accept "dog (собака)" style round-tripped labels too
    code = text.split(" ")[0]
    if code in _PET_KIND_ALIASES:
        return _PET_KIND_ALIASES[code]
    if text in _PET_KIND_ALIASES:
        return _PET_KIND_ALIASES[text]
    raise ValueError(f"неизвестный вид питомца: {value!r}")


def _parse_size(value: object) -> DogSizeClass | None:
    text = _cell_str(value).lower()
    if not text:
        return None
    code = text.split(" ")[0]
    if code in _SIZE_ALIASES:
        return _SIZE_ALIASES[code]
    if text in _SIZE_ALIASES:
        return _SIZE_ALIASES[text]
    raise ValueError(f"неизвестный размер: {value!r}")


def _parse_int(value: object) -> int | None:
    text = _cell_str(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError(f"не число: {value!r}") from exc


def known_tariff_kinds(catalog: Catalog) -> list[str]:
    seen = {a.tariff_kind for a in catalog.accommodations} | {
        r.accommodation_tariff for r in catalog.daily_rates
    }
    return sorted(seen)


def build_prices_workbook(catalog: Catalog) -> Workbook:
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = SHEET_PRICES
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    for rate in catalog.daily_rates:
        ws.append(
            [
                rate.accommodation_tariff,
                PET_KIND_LABELS.get(rate.pet_kind, rate.pet_kind.value),
                SIZE_LABELS.get(rate.size_class, "") if rate.size_class else "",
                "да" if rate.is_young else "нет",
                "да" if rate.under_one_month else "нет",
                rate.price_rub if rate.price_rub is not None else "",
                rate.price_from_rub if rate.price_from_rub is not None else "",
                rate.price_to_rub if rate.price_to_rub is not None else "",
                rate.price_rub_season if rate.price_rub_season is not None else "",
                rate.price_from_rub_season if rate.price_from_rub_season is not None else "",
                rate.price_to_rub_season if rate.price_to_rub_season is not None else "",
                "да" if rate.unavailable else "нет",
            ]
        )
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
            14, len(header) + 2
        )

    legend = wb.create_sheet(SHEET_LEGEND)
    legend.append(["Как редактировать", ""])
    legend.append([
        "Чтобы изменить цену — впишите новое число в нужную ячейку на листе «Цены».",
        "",
    ])
    legend.append([
        "Чтобы добавить новую позицию — допишите новую строку снизу со всеми колонками.",
        "",
    ])
    legend.append([
        "Чтобы убрать позицию из прайса — удалите всю строку (не просто очистите цену).",
        "",
    ])
    legend.append(["", ""])
    legend.append(["Код тарифа (колонка «Тариф»)", "Используется сейчас"])
    for tariff in known_tariff_kinds(catalog):
        legend.append([tariff, ""])
    legend.append(["", ""])
    legend.append(["Вид питомца (колонка «Вид питомца»)", ""])
    for kind, label in PET_KIND_LABELS.items():
        legend.append([kind.value, label])
    legend.append(["", ""])
    legend.append(["Размер — только для собак (колонка «Размер»)", ""])
    for size, label in SIZE_LABELS.items():
        legend.append([size.value, label])
    legend.append(["", ""])
    legend.append([
        "«Недоступно» = да значит, что этот тариф для этого вида/размера не предлагается "
        "(цену можно не заполнять).",
        "",
    ])
    legend.column_dimensions["A"].width = 70
    legend.column_dimensions["B"].width = 30
    return wb


def export_prices(catalog: Catalog, out_path: Path) -> Path:
    wb = build_prices_workbook(catalog)
    wb.save(out_path)
    return out_path


class PriceImportError(Exception):
    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        super().__init__("; ".join(errors))


def parse_prices_workbook(path: Path) -> list[DailyRate]:
    wb = load_workbook(path, data_only=True)
    if SHEET_PRICES not in wb.sheetnames:
        raise PriceImportError([f"В файле нет листа «{SHEET_PRICES}»."])
    ws = wb[SHEET_PRICES]

    rates: list[DailyRate] = []
    errors: list[str] = []
    seen_keys: set[tuple] = set()

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = list(row) + [None] * (len(HEADERS) - len(row))
        (
            tariff_raw, kind_raw, size_raw, young_raw, under_month_raw,
            price_raw, from_raw, to_raw,
            season_price_raw, season_from_raw, season_to_raw,
            unavailable_raw,
        ) = cells[: len(HEADERS)]

        tariff = _cell_str(tariff_raw)
        if not tariff:
            errors.append(f"Строка {row_idx}: не указан тариф.")
            continue
        try:
            pet_kind = _parse_pet_kind(kind_raw)
            size_class = _parse_size(size_raw)
            is_young = _parse_bool(young_raw)
            under_one_month = _parse_bool(under_month_raw)
            price_rub = _parse_int(price_raw)
            price_from_rub = _parse_int(from_raw)
            price_to_rub = _parse_int(to_raw)
            price_rub_season = _parse_int(season_price_raw)
            price_from_rub_season = _parse_int(season_from_raw)
            price_to_rub_season = _parse_int(season_to_raw)
            unavailable = _parse_bool(unavailable_raw)
        except ValueError as exc:
            errors.append(f"Строка {row_idx}: {exc}")
            continue

        if pet_kind == PetKind.DOG and size_class is None:
            errors.append(f"Строка {row_idx}: для собак нужно указать размер.")
            continue
        if pet_kind != PetKind.DOG and size_class is not None:
            errors.append(f"Строка {row_idx}: размер задаётся только для собак.")
            continue
        if not unavailable and price_rub is None and price_from_rub is None:
            errors.append(
                f"Строка {row_idx}: не задана цена (или отметьте «Недоступно» = да)."
            )
            continue

        key = (tariff, pet_kind, size_class, is_young, under_one_month)
        if key in seen_keys:
            errors.append(f"Строка {row_idx}: дублирует уже встреченную комбинацию {key}.")
            continue
        seen_keys.add(key)

        rates.append(
            DailyRate(
                accommodation_tariff=tariff,
                pet_kind=pet_kind,
                size_class=size_class,
                is_young=is_young,
                under_one_month=under_one_month,
                price_rub=price_rub,
                price_from_rub=price_from_rub,
                price_to_rub=price_to_rub,
                price_rub_season=price_rub_season,
                price_from_rub_season=price_from_rub_season,
                price_to_rub_season=price_to_rub_season,
                unavailable=unavailable,
            )
        )

    if errors:
        raise PriceImportError(errors)
    if not rates:
        raise PriceImportError(["В файле не нашлось ни одной валидной строки с ценой."])
    return rates


def _rate_key(rate: DailyRate) -> tuple:
    return (
        rate.accommodation_tariff,
        rate.pet_kind,
        rate.size_class,
        rate.is_young,
        rate.under_one_month,
    )


def diff_summary(old: list[DailyRate], new: list[DailyRate]) -> str:
    old_by_key = {_rate_key(r): r for r in old}
    new_by_key = {_rate_key(r): r for r in new}
    added = [k for k in new_by_key if k not in old_by_key]
    removed = [k for k in old_by_key if k not in new_by_key]
    changed = [
        k
        for k in new_by_key
        if k in old_by_key and old_by_key[k].model_dump() != new_by_key[k].model_dump()
    ]
    parts = []
    if added:
        parts.append(f"добавлено: {len(added)}")
    if changed:
        parts.append(f"изменено: {len(changed)}")
    if removed:
        parts.append(f"удалено: {len(removed)}")
    return ", ".join(parts) if parts else "без изменений"


_HEADER_COMMENT = (
    "# Цены на размещение (матрица тариф × вид питомца × размер × возраст × сезон).\n"
    "# Этот файл перезаписывается полностью при импорте через /prices в боте.\n"
    "# unavailable: true — тариф для этого вида/размера питомца не предлагается.\n"
    "# price_rub_season / price_from_rub_season / price_to_rub_season — сезон\n"
    "# отпусков и праздников; применяется, когда заезд попадает в период из\n"
    "# price_rules.yaml: season_periods.\n\n"
)


def save_daily_rates(config_dir: Path, rates: list[DailyRate]) -> Path:
    path = config_dir / "daily_rates.yaml"
    data = {
        "daily_rates": [r.model_dump(mode="json", exclude_none=True) for r in rates]
    }
    # keep unavailable only when true, to match the hand-written style
    for entry, rate in zip(data["daily_rates"], rates):
        if not rate.unavailable:
            entry.pop("unavailable", None)
    body = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=False)
    path.write_text(_HEADER_COMMENT + body, encoding="utf-8")
    return path

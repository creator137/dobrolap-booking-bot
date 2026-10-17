from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from dobrolap_bot.config.loader import load_catalog
from dobrolap_bot.services import price_admin


@pytest.fixture(scope="module")
def catalog():
    root = Path(__file__).resolve().parents[1]
    return load_catalog(root / "config")


def test_export_import_round_trip_no_changes(catalog, tmp_path):
    out = tmp_path / "prices.xlsx"
    price_admin.export_prices(catalog, out)

    rates = price_admin.parse_prices_workbook(out)

    assert len(rates) == len(catalog.daily_rates)
    assert price_admin.diff_summary(catalog.daily_rates, rates) == "без изменений"


def test_edit_price_and_add_row(catalog, tmp_path):
    out = tmp_path / "prices.xlsx"
    price_admin.export_prices(catalog, out)

    wb = load_workbook(out)
    ws = wb[price_admin.SHEET_PRICES]
    ws.cell(row=2, column=6).value = 9999  # change the first price cell
    ws.append(["brand_new_tariff", "cat", "", "нет", "нет", 500, "", "", "", "", "", "нет"])
    wb.save(out)

    rates = price_admin.parse_prices_workbook(out)
    assert len(rates) == len(catalog.daily_rates) + 1

    summary = price_admin.diff_summary(catalog.daily_rates, rates)
    assert "добавлено: 1" in summary
    assert "изменено: 1" in summary

    new = next(r for r in rates if r.accommodation_tariff == "brand_new_tariff")
    assert new.price_rub == 500


def test_remove_row(catalog, tmp_path):
    out = tmp_path / "prices.xlsx"
    price_admin.export_prices(catalog, out)

    wb = load_workbook(out)
    ws = wb[price_admin.SHEET_PRICES]
    ws.delete_rows(2)
    wb.save(out)

    rates = price_admin.parse_prices_workbook(out)
    assert len(rates) == len(catalog.daily_rates) - 1
    assert "удалено: 1" in price_admin.diff_summary(catalog.daily_rates, rates)


def test_invalid_price_raises_with_row_number(catalog, tmp_path):
    out = tmp_path / "prices.xlsx"
    price_admin.export_prices(catalog, out)

    wb = load_workbook(out)
    ws = wb[price_admin.SHEET_PRICES]
    ws.cell(row=2, column=6).value = "много денег"
    wb.save(out)

    with pytest.raises(price_admin.PriceImportError) as exc_info:
        price_admin.parse_prices_workbook(out)

    assert "Строка 2" in exc_info.value.errors[0]


def test_dog_requires_size_class(catalog, tmp_path):
    out = tmp_path / "prices.xlsx"
    price_admin.export_prices(catalog, out)

    wb = load_workbook(out)
    ws = wb[price_admin.SHEET_PRICES]
    ws.append(["vip_plus", "dog", "", "нет", "нет", 1000, "", "", "", "", "", "нет"])
    wb.save(out)

    with pytest.raises(price_admin.PriceImportError) as exc_info:
        price_admin.parse_prices_workbook(out)

    assert any("размер" in e for e in exc_info.value.errors)


def test_save_daily_rates_round_trips_through_loader(catalog, tmp_path):
    rates = list(catalog.daily_rates)
    price_admin.save_daily_rates(tmp_path, rates)

    import yaml

    data = yaml.safe_load((tmp_path / "daily_rates.yaml").read_text(encoding="utf-8"))
    assert len(data["daily_rates"]) == len(rates)
